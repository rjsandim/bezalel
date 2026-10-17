"""Ensaios de exportação local com dados sintéticos, sem rede ou produto."""
from copy import deepcopy
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from openpyxl import load_workbook

SKILL = Path(__file__).resolve().parents[1]
SCRIPT = SKILL / 'scripts/build_test_workbook.py'
spec = importlib.util.spec_from_file_location('workbook_builder', SCRIPT)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
FIXTURE = json.loads((SKILL / 'tests/fixtures/synthetic.json').read_text())


class WorkbookTests(unittest.TestCase):
    def roundtrip(self, data):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'fixture.xlsx'
            builder.build_workbook(data).save(output)
            return load_workbook(output)

    def test_empty_template_keeps_six_sheets_and_no_claimed_results(self):
        w = self.roundtrip({})
        self.assertEqual(w.sheetnames, ['Resumo', 'Guia de execução', 'Casos de teste', 'Dados do ambiente', 'Consultas SQL', 'Backlog'])
        for name, col in [('Casos de teste', 2), ('Dados do ambiente', 1), ('Consultas SQL', 1), ('Backlog', 1)]:
            self.assertFalse(any(w[name].cell(row, col).value for row in range(6, w[name].max_row + 1)))
        self.assertEqual(w['Casos de teste'].freeze_panes, 'C6')
        self.assertTrue(w.calculation.fullCalcOnLoad)

    def test_results_dates_links_and_formula_fields_survive_export(self):
        w = self.roundtrip(FIXTURE)
        cases = w['Casos de teste']
        self.assertEqual(cases['O6'].value, 'Falhou')
        self.assertEqual(cases['T6'].value, 'B01')
        self.assertEqual(cases['Q6'].value.date().isoformat(), '2026-01-01')
        self.assertEqual(w['Backlog']['B6'].value, 'T001')
        self.assertEqual(cases['S7'].value, '=IF(OR(O7="Falhou",O7="Bloqueado"),"Criar item","")')
        self.assertEqual(w['Resumo']['B6'].value, "=COUNTA('Casos de teste'!$B$6:$B$205)")
        self.assertIn('<>Fechado', w['Resumo']['B14'].value)
        self.assertIn('<>Descartado', w['Resumo']['B14'].value)

    def test_capacity_expands_formulas_filters_validation_and_coverage(self):
        data = {'cases': [{'id': f'T{i:03}'} for i in range(1, 204)],
                'backlog': [{'id': f'B{i}', 'testId': 'T001'} for i in range(102)],
                'coverage': [{'dimension': f'Fictícia {i}'} for i in range(5)]}
        w = self.roundtrip(data)
        for name, last, column in [('Casos de teste', 208, 'T'), ('Backlog', 107, 'P')]:
            sheet = w[name]
            table = next(iter(sheet.tables.values()))
            self.assertEqual(table.ref, f'A5:{column}{last}')
            self.assertEqual(table.autoFilter.ref, table.ref)
            self.assertTrue(all(next(iter(v.sqref.ranges)).max_row == last for v in sheet.data_validations.dataValidation))
            self.assertTrue(all(next(iter(c.sqref.ranges)).max_row == last for c in sheet.conditional_formatting))
        self.assertIn('O208', w['Casos de teste']['S208'].value)
        self.assertIn('$208', w['Resumo']['B6'].value)
        self.assertIn('$107', w['Resumo']['B14'].value)
        self.assertEqual(w['Resumo']['A22'].value, 'Fictícia 4')
        self.assertIn('A24:H24', {str(r) for r in w['Resumo'].merged_cells.ranges})

    def test_untrusted_formula_like_content_stays_literal_text(self):
        data = deepcopy(FIXTURE)
        data['cases'][0]['steps'] = '=HYPERLINK("https://example.invalid","fixture")'
        w = self.roundtrip(data)
        self.assertEqual(w['Casos de teste']['J6'].data_type, 's')
        self.assertEqual(w['Casos de teste']['J6'].value, data['cases'][0]['steps'])
        self.assertFalse(w._external_links)

    def test_invalid_ids_references_status_and_dates_are_rejected(self):
        changes = [('id', 'wrong'), ('sqlIds', 'Q99'), ('backlogId', 'missing'), ('result', 'OK'), ('date', '2026-02-30')]
        for field, value in changes:
            with self.subTest(field=field):
                data = deepcopy(FIXTURE); data['cases'][0][field] = value
                with self.assertRaises(ValueError):
                    builder.build_workbook(data)
        with self.assertRaises(ValueError):
            builder.build_workbook({'cases': [{'id': 'T001'}, {'id': 'T001'}]})

    def test_standalone_copy_runs_and_refuses_overwrite(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            isolated = root / 'skill'
            shutil.copytree(SKILL, isolated)
            output = root / 'result.xlsx'
            args = [sys.executable, '-B', str(isolated / 'scripts/build_test_workbook.py'), '-', str(output)]
            result = subprocess.run(args, cwd=root, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            before = output.read_bytes()
            result = subprocess.run(args, cwd=root, capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(output.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
