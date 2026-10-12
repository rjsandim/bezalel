#!/usr/bin/env python3
"""Gera um roteiro Excel local a partir de JSON, sem executar casos ou consultas."""
import argparse
from copy import copy
from datetime import date
import json
from pathlib import Path
import re
import sys

try:
    from openpyxl import load_workbook
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
except ModuleNotFoundError:
    raise SystemExit("Dependência ausente: use um ambiente com openpyxl 3.1.x disponível.") from None

TEMPLATE = Path(__file__).resolve().parents[1] / 'assets/test-plan-template.xlsx'
FIELDS = {
    'guide': ('Guia de execução', 'section item guidance link required expected'.split(), 9),
    'cases': ('Casos de teste', 'order id priority area integration unit channel fixture precondition steps expectedSurface expectedDatabase sqlIds evidence result observed date executor _action backlogId'.split(), 200),
    'environmentData': ('Dados do ambiente', 'code use unit integration primaryEntity relatedEntities initialState expectedState ids entryLink databaseSource notes'.split(), 20),
    'queries': ('Consultas SQL', 'id when objective sql expectedBefore expectedAfter care'.split(), 50),
    'backlog': ('Backlog', 'id testId date unit integration severity type summary expected observed evidence sqlIds owner status issue retest'.split(), 100),
}
COVERAGE = 'dimension uiApi rules dataSetup success failureRetry migrationCompatibility security'.split()
PRIORITIES = {'Crítica', 'Alta', 'Média', 'Baixa'}
RESULTS = {'Não iniciado', 'Em execução', 'Aprovado', 'Falhou', 'Bloqueado'}
BACKLOG_STATES = {'Novo', 'Em análise', 'Em correção', 'Pronto para reteste', 'Fechado', 'Descartado'}


def validate_spec(spec):
    if not isinstance(spec, dict):
        raise ValueError('A especificação deve ser um objeto JSON.')
    for key in [*FIELDS, 'coverage']:
        rows = spec.get(key, [])
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError(f'{key}: use uma lista de objetos.')
        for row in rows:
            if any(isinstance(value, (dict, list)) for value in row.values()):
                raise ValueError(f'{key}: campos devem conter valores simples.')
            if row.get('date'):
                date.fromisoformat(str(row['date']))
    ids = {}
    for key, pattern in [('cases', r'T\d{3,}'), ('queries', r'Q\d+'), ('backlog', r'\S+')]:
        values = [row.get('id') for row in spec.get(key, [])]
        if any(not isinstance(v, str) or not re.fullmatch(pattern, v) for v in values) or len(set(values)) != len(values):
            raise ValueError(f'{key}: IDs ausentes, inválidos ou duplicados.')
        ids[key] = set(values)
    for key, field, allowed in [('cases', 'result', RESULTS), ('cases', 'priority', PRIORITIES), ('backlog', 'severity', PRIORITIES), ('backlog', 'status', BACKLOG_STATES), ('backlog', 'retest', {'Não retestado', 'Aprovado', 'Falhou', 'Bloqueado'})]:
        if any(row.get(field) and row[field] not in allowed for row in spec.get(key, [])):
            raise ValueError(f'{key}: valor inválido em {field}.')
    for key in ('cases', 'backlog'):
        for row in spec.get(key, []):
            refs = re.split(r'[,;\s]+', str(row.get('sqlIds') or '').strip())
            if any(ref and ref not in ids['queries'] for ref in refs):
                raise ValueError(f'{key}: referência SQL inexistente.')
            field, target = ('backlogId', 'backlog') if key == 'cases' else ('testId', 'cases')
            if row.get(field) and row[field] not in ids[target]:
                raise ValueError(f'{key}: referência inexistente em {field}.')


def write_value(cell, value):
    cell.value = value
    if isinstance(value, str):
        # Entrada é conteúdo: somente fórmulas internas do gerador são executáveis.
        cell.data_type = 's'


def fill_sheet(sheet, rows, fields, minimum):
    end = 5 + max(minimum, len(rows))
    old_end = sheet.max_row
    for row_number in range(6, end + 1):
        item = rows[row_number - 6] if row_number - 6 < len(rows) else None
        for col, field in enumerate(fields, 1):
            cell = sheet.cell(row_number, col)
            if row_number > old_end:
                cell._style = copy(sheet.cell(6 + (row_number % 2), col)._style)
            value = item.get(field) if item else None
            if item and field == 'order' and value is None:
                value = row_number - 5
            if item and field == 'result' and not value:
                value = 'Não iniciado'
            if field == 'date' and value:
                value = date.fromisoformat(str(value))
            write_value(cell, value)
        if sheet.title == 'Casos de teste':
            sheet.cell(row_number, 19, f'=IF(OR(O{row_number}="Falhou",O{row_number}="Bloqueado"),"Criar item","")')
        sheet.row_dimensions[row_number].height = 64 if item else 22
    last_col = get_column_letter(len(fields))
    for table in sheet.tables.values():
        table.ref = f'A5:{last_col}{end}'
        if table.autoFilter:
            table.autoFilter.ref = table.ref
    for validation in sheet.data_validations.dataValidation:
        col = get_column_letter(next(iter(validation.sqref.ranges)).min_col)
        validation.sqref = f'{col}6:{col}{end}'
    formats = ConditionalFormattingList()
    for conditional in sheet.conditional_formatting:
        for area in conditional.sqref.ranges:
            col = get_column_letter(area.min_col)
            for rule in sheet.conditional_formatting[conditional]:
                formats.add(f'{col}6:{col}{end}', copy(rule))
    sheet.conditional_formatting = formats
    return end


def build_workbook(spec):
    validate_spec(spec)
    workbook = load_workbook(TEMPLATE)
    workbook.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    ends = {}
    for key, (name, fields, minimum) in FIELDS.items():
        sheet = workbook[name]
        rows = spec.get(key, [])
        if key == 'guide' and not rows:
            rows = [dict(zip(fields, values)) for values in sheet.iter_rows(min_row=6, values_only=True)]
        ends[key] = fill_sheet(sheet, rows, fields, minimum)
        sheet.freeze_panes = 'C6' if key in ('cases', 'backlog') else 'A6'
    for cell, key in [('A1', 'title'), ('A2', 'subtitle'), ('A3', 'warning')]:
        if key in spec:
            write_value(workbook['Casos de teste'][cell], spec[key])
    if 'warning' in spec:
        write_value(workbook['Guia de execução']['A3'], spec['warning'])
    summary = workbook['Resumo']
    summary.freeze_panes = 'A6'
    for row in range(6, 15):
        cell = summary.cell(row, 2)
        cell.value = cell.value.replace('$205', f'${ends["cases"]}').replace('$105', f'${ends["backlog"]}')
    coverage = spec.get('coverage', [])
    acceptance = summary['A20'].value
    acceptance_style = copy(summary['A20']._style)
    summary.unmerge_cells('A20:H20')
    summary['A20'] = None
    for index in range(max(1, len(coverage))):
        for col, field in enumerate(COVERAGE, 1):
            cell = summary.cell(18 + index, col)
            cell._style = copy(summary.cell(18, col)._style)
            write_value(cell, coverage[index].get(field) if index < len(coverage) else None)
        summary.row_dimensions[18 + index].height = 44
    row = 19 + max(1, len(coverage))
    summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    write_value(summary.cell(row, 1), 'Critério de aceite geral: ' + str(spec['generalAcceptance']) if 'generalAcceptance' in spec else acceptance)
    summary.cell(row, 1)._style = acceptance_style
    summary.row_dimensions[row].height = 44
    return workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec', help='JSON ou - para template vazio')
    parser.add_argument('output', type=Path, help='Novo arquivo .xlsx; nunca sobrescreve')
    args = parser.parse_args()
    try:
        if args.output.suffix.lower() != '.xlsx':
            raise ValueError('A saída deve ter extensão .xlsx.')
        spec = {} if args.spec == '-' else json.loads(Path(args.spec).read_text(encoding='utf-8'))
        workbook = build_workbook(spec)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open('xb') as stream:
            workbook.save(stream)
        print('Workbook criado; fórmulas serão recalculadas ao abrir no Excel/LibreOffice. Nenhum teste executado.')
        return 0
    except (ValueError, OSError) as error:
        # Não ecoar conteúdo/valores privados de entrada.
        print(f'Não foi possível gerar o workbook ({type(error).__name__}); confira JSON, referências e destino novo.', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
